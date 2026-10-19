"""Build the student Word report with bundled python-docx and real artifacts."""
from pathlib import Path
import hashlib
import json
import re
from urllib.parse import quote
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'reports/实验一报告_郑航宇_2024280087.docx'
SCENARIO = {'continuous': '连续满载', 'staggered_burst': '错峰突发',
            'synchronized_burst': '同步突发', 'light_random': '轻载随机'}
MODE = {'circuit': '电路', 'packet': '分组'}
REPOSITORY = 'https://github.com/erjimeidianle/modern-switching-experiment-01/blob/main/'


def font(run, size=None, bold=None, mono=False):
    run.font.name = 'Consolas' if mono else 'Times New Roman'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'SimSun')
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    return run


def inline(paragraph, text, size=None):
    tokens = re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))', text)
    for token in tokens:
        if not token:
            continue
        link = re.fullmatch(r'\[([^]]+)\]\(([^)]+)\)', token)
        if link:
            label, target = link.groups()
            address = target if '://' in target else REPOSITORY + quote(target, safe='/')
            relation = paragraph.part.relate_to(
                address, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                is_external=True)
            element = OxmlElement('w:hyperlink')
            element.set(qn('r:id'), relation)
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '215A85')
            props.append(color)
            run.append(props)
            value = OxmlElement('w:t'); value.text = label
            run.append(value); element.append(run); paragraph._p.append(element)
        elif token.startswith('**'):
            font(paragraph.add_run(token[2:-2]), size, True)
        elif token.startswith('`'):
            font(paragraph.add_run(token[1:-1]), min(size or 11, 9.5), mono=True)
        else:
            font(paragraph.add_run(token), size)


def table(doc, lines, number):
    values = [[c.strip() for c in line.strip().strip('|').split('|')] for line in lines
              if not re.match(r'^\|[ :|\-]+\|$', line)]
    columns = len(values[0])
    widths = {2: [5.2, 12.4], 4: [1.2, 5.8, 6.6, 4.0],
              6: [2.6, 1.8, 3.1, 3.2, 3.0, 3.9],
              8: [2.5, 1.2, 2.4, 2.4, 2.3, 2.3, 2.0, 2.5]}[columns]
    # Fit all tables to the usable seven-inch page width.
    widths = [v * 17.78 / sum(widths) for v in widths]
    labels = {1: '指标定义与计量单位', 2: '两用户手工推演', 3: '默认配置下的四场景指标',
              4: '原始有效载荷守恒核查', 5: '包头参数变化下的分组指标'}
    caption = doc.add_paragraph(f'表 {number}  {labels[number]}')
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(5)
    for run in caption.runs:
        font(run, 10, True)
    tab = doc.add_table(rows=0, cols=columns)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for col, width in zip(tab.columns, widths):
        col.width = Cm(width)
    props = tab._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + side)
        for attr, value in (('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')):
            edge.set(qn('w:' + attr), value)
        borders.append(edge)
    props.append(borders)
    for row_index, row in enumerate(values):
        cells = tab.add_row().cells
        trprops = tab.rows[-1]._tr.get_or_add_trPr()
        trprops.append(OxmlElement('w:cantSplit'))
        if row_index == 0:
            trprops.append(OxmlElement('w:tblHeader'))
        for index, (cell, text) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cellprops = cell._tc.get_or_add_tcPr()
            padding = OxmlElement('w:tcMar')
            for edge, value in (('top', '75'), ('bottom', '75'), ('left', '90'), ('right', '90')):
                part = OxmlElement('w:' + edge)
                part.set(qn('w:w'), value); part.set(qn('w:type'), 'dxa')
                padding.append(part)
            cellprops.append(padding)
            if row_index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EEF3')
                cellprops.append(shade)
                replacements = {'有效利用率%': '有效利用率\n%', '线路忙碌率%': '线路忙碌率\n%',
                                '平均时延ms': '平均时延\nms', '最大时延ms': '最大时延\nms',
                                '丢失率%': '丢失率\n%', '积压bit': '积压\nbit'}
                text = replacements.get(text, text)
            else:
                text = SCENARIO.get(text, MODE.get(text, text))
                if number in (3, 5) and re.fullmatch(r'\d+\.\d+', text):
                    text = f'{float(text):.4f}'.rstrip('0').rstrip('.')
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            paragraph.paragraph_format.keep_with_next = False
            paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT if columns == 2 or
                                   (number == 2 and index in (1, 2, 3)) else WD_ALIGN_PARAGRAPH.CENTER)
            inline(paragraph, text, 9.5 if columns >= 6 else 10)
            if row_index == 0:
                for run in paragraph.runs:
                    font(run, 10, True)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.line_spacing = 0.3


def main():
    source = (ROOT / 'REPORT.md').read_text(encoding='utf-8')
    baseline = json.loads((ROOT / 'results/comparison.json').read_text(encoding='utf-8'))
    # Validate the numerical source before repackaging it.
    for row in baseline['rows']:
        key = f"{row['scenario']}/{row['mode']}"
        raw = baseline['traces'][key]
        assert raw['generated_payload_bits'] == sum(raw[k] for k in (
            'delivered_payload_bits', 'dropped_payload_bits', 'backlog_payload_bits'))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.75)
    sec.footer_distance = Inches(.3)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'SimSun')
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(style._element.findall('.//' + qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in (('Title', 21), ('Heading 1', 14), ('Heading 2', 12)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(6)
        doc.styles[name].paragraph_format.keep_with_next = True
    title = doc.add_paragraph('电路交换与分组交换资源利用率比较', style='Title')
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    for text in ('现代交换原理实验一', '郑航宇  2024280087  GitHub erjimeidianle', '实验日期 2026年10月19日'):
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(4)
    summary = doc.add_paragraph()
    inline(summary, '实验结果表明，两种交换方式的优劣取决于业务负载。连续满载下电路有效利用率为100%，分组为79.6875%；错峰热点下分组有效利用率为75%、平均完成时延为1 ms，优于本次配置下的电路模型。下文给出模型、手算、真实数据、反例测试和包头参数实验。')

    # Convert the existing report content, preserving narrative and evidence.
    lines = source[source.index('## 实验目的'):].splitlines()
    index = 0; code = False; table_number = 0; figure_number = 0
    while index < len(lines):
        line = lines[index].strip(); index += 1
        if not line:
            continue
        if line.startswith('```'):
            code = not code
            continue
        if code:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.08
            font(paragraph.add_run(line), 8.5, mono=True)
        elif line.startswith('|'):
            block = [line]
            while index < len(lines) and lines[index].strip().startswith('|'):
                block.append(lines[index].strip()); index += 1
            table_number += 1
            table(doc, block, table_number)
        elif (match := re.fullmatch(r'!\[([^]]+)\]\(([^)]+)\)', line)):
            label, target = match.groups(); figure_number += 1
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(3)
            picture = paragraph.add_run().add_picture(str(ROOT / target), width=Inches(7))
            picture._inline.docPr.set('descr', label)
            caption = doc.add_paragraph(f'图 {figure_number}  {label}')
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(8)
            for run in caption.runs:
                font(run, 9.5)
        elif line.startswith('### '):
            doc.add_heading(re.sub(r'[：:、。]', ' ', line[4:]), level=2)
        elif line.startswith('## '):
            doc.add_heading(re.sub(r'[：:、。]', ' ', line[3:]), level=1)
        else:
            paragraph = doc.add_paragraph()
            inline(paragraph, line)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('第 '), 9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    font(footer.add_run(' 页'), 9)
    doc.core_properties.title = '电路交换与分组交换资源利用率比较'
    doc.core_properties.author = '郑航宇'
    doc.core_properties.subject = '现代交换原理实验一'
    doc.core_properties.keywords = '电路交换 分组交换 FIFO 有效利用率 完成时延'
    OUTPUT.parent.mkdir(exist_ok=True)
    doc.save(OUTPUT)
    with ZipFile(OUTPUT) as package:
        assert package.testzip() is None
        assert len([n for n in package.namelist() if n.startswith('word/media/')]) == 4
    reopened = Document(OUTPUT)
    assert len(reopened.tables) == 5
    assert [len(t.rows) for t in reopened.tables] == [7, 3, 9, 9, 13]
    body = '\n'.join(p.text for p in reopened.paragraphs)
    assert all(s in body for s in ('郑航宇', '2024280087', '70/70', '24 passed', '37.5%'))
    evidence = {'status': 'created', 'output': str(OUTPUT), 'tables': 5, 'embedded_figures': 4,
                'source_sha256': hashlib.sha256(source.encode('utf-8')).hexdigest(),
                'docx_sha256': hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
                'rounding': 'Metric tables display at most four decimal places; raw JSON and CSV retained.'}
    (ROOT / '.build/word_report').mkdir(exist_ok=True)
    (ROOT / '.build/word_report/build_validation.json').write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(evidence, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
